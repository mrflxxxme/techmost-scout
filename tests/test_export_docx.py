"""Проверки выгрузки отчёта в DOCX (Э4-И5, SPEC Э4-2 часть 2).

`export_docx` принимает уже готовый `ExportView` — тем же контрактом, что и
страница отчёта, поэтому в тестах он собирается напрямую, без базы данных
(инвариант №7: сети и файловой БД тест не требует, только `tmp_path`).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

from app.export.docx import DRAFT_NOTICE, export_docx
from app.export.markdown import REVIEWED_FOOTER
from app.export.view import ExportView
from app.schemas import (
    Evidence,
    GovSupportMatch,
    MaturityStage,
    Report,
    ReportSummary,
    RequirementCheck,
    SolutionCard,
)


def _make_card(**overrides) -> SolutionCard:
    fields = {
        "title": "Регенерация катализатора добавками молибдена",
        "essence": "Продление срока службы катализатора за счёт присадки.",
        "applicability": "Закрывает задачу продления межрегенерационного цикла.",
        "next_step": "Запросить образец у разработчика.",
        "evidence": [Evidence(claim="Испытан на действующей установке", url="https://example.org/proof")],
    }
    fields.update(overrides)
    return SolutionCard(**fields)


def _make_report(cards: list[SolutionCard], contest_cards: list[SolutionCard] | None = None, **overrides) -> Report:
    fields = {
        "request_id": 1,
        "brief": overrides.pop("brief"),
        "cards": cards,
        "contest_cards": contest_cards or [],
        "summary": ReportSummary(recommendation="Начать с варианта №1."),
        "generated_at": datetime(2026, 8, 27, 9, 0),
    }
    fields.update(overrides)
    return Report(**fields)


def _make_view(report: Report, *, is_draft: bool = True, reviewer_name: str = "", validated_at=None) -> ExportView:
    return ExportView(
        request_id=report.request_id,
        report=report,
        card_ids={},
        is_draft=is_draft,
        reviewer_name=reviewer_name,
        validated_at=validated_at,
    )


def _full_text(path: Path) -> str:
    """Весь текст документа одной строкой — абзацы и ячейки таблиц."""
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_файл_создаётся_с_именем_выгрузки_и_читается_обратно(tmp_path: Path, brief):
    """Критерий 1: имя `report-<id>-<дата>.docx`, файл открывается python-docx,
    в тексте есть заголовок карточки и URL подтверждения."""
    card = _make_card(title="Катализатор с молибденовой присадкой")
    report = _make_report([card], brief=brief)
    view = _make_view(report, is_draft=True)
    out_path = tmp_path / "report-1-2026-08-27.docx"

    result = export_docx(view, out_path)

    assert result == out_path
    assert out_path.exists()
    text = _full_text(out_path)
    assert "Катализатор с молибденовой присадкой" in text
    assert "https://example.org/proof" in text


def test_черновик_есть_пометка_и_нет_проверено(tmp_path: Path, brief):
    """Критерий 2 (часть 1): `is_draft` → есть DRAFT_NOTICE, нет «проверено:»."""
    report = _make_report([_make_card()], brief=brief)
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    assert DRAFT_NOTICE in text
    assert "проверено:" not in text


def test_validated_ровно_подпись_и_ни_одного_черновика(tmp_path: Path, brief):
    """Критерий 2 (часть 2): при `validated` — ровно строка РЕШ-197, ни одного
    вхождения «ЧЕРНОВИК»."""
    report = _make_report([_make_card()], brief=brief)
    when = datetime(2026, 8, 27)
    view = _make_view(report, is_draft=False, reviewer_name="Иван Иванов", validated_at=when)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    expected = REVIEWED_FOOTER.format(name="Иван Иванов", date="27.08.2026")
    assert expected in text
    assert "ЧЕРНОВИК" not in text


def test_карточка_без_evidence_не_попадает_в_docx(tmp_path: Path, brief):
    """РЕШ-204: карточка без единого подтверждения не печатается вовсе."""
    bare = _make_card(title="Карточка без подтверждений", maturity_stage=MaturityStage.LAB, trl_estimate=5, evidence=[])
    report = _make_report([bare], brief=brief)
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    assert "Карточка без подтверждений" not in text
    assert "Стадия зрелости" not in text
    assert "TRL" not in text
    assert "Пригодных вариантов не найдено." in text


def test_карточка_с_evidence_печатает_стадию_и_trl(tmp_path: Path, brief):
    """Регрессия к инварианту №1: подтверждённая карточка стадию всё же печатает."""
    card = _make_card(
        maturity_stage=MaturityStage.PILOT,
        trl_estimate=6,
        evidence=[Evidence(claim="Пилот на объекте", url="https://example.org/pilot")],
    )
    report = _make_report([card], brief=brief)
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    assert "прошла пилот на объекте" in text
    assert "TRL ≈ 6" in text


def test_карточка_без_evidence_пропускается_а_порядок_остальных_сохраняется(tmp_path: Path, brief):
    """Смешанный список: карточка без evidence выпадает, соседние сохраняют порядок."""
    first = _make_card(title="Первая с подтверждением")
    bare = _make_card(title="Вторая без подтверждения", evidence=[])
    third = _make_card(title="Третья с подтверждением")
    report = _make_report([first, bare, third], brief=brief)
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    assert "Вторая без подтверждения" not in text
    assert text.index("Первая с подтверждением") < text.index("Третья с подтверждением")


def test_конкурсная_карточка_без_evidence_не_попадает_и_секция_скрывается(tmp_path: Path, brief):
    """Если у всех конкурсных карточек пустой evidence — заголовок секции не печатается."""
    main_card = _make_card(title="Основная карточка")
    bare_contest = _make_card(title="Конкурсная без подтверждения", evidence=[])
    report = _make_report([main_card], contest_cards=[bare_contest], brief=brief)
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    assert "Конкурсная без подтверждения" not in text
    assert "Ранние разработки и команды" not in text


def test_секция_сравнения_печатается_таблицей(tmp_path: Path, brief):
    """Сравнение вариантов идёт таблицей: заголовки и значения строк видны."""
    report = _make_report(
        [_make_card()],
        brief=brief,
        summary=ReportSummary(
            recommendation="Начать с варианта №1.",
            comparison_rows=[{"Вариант": "Молибден", "Срок": "24 месяца"}],
        ),
    )
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    document = Document(str(tmp_path / "report-1-2026-08-27.docx"))
    assert len(document.tables) == 1
    header = [cell.text for cell in document.tables[0].rows[0].cells]
    assert header == ["Вариант", "Срок"]


def test_конкурсная_секция_и_разработчик_печатаются(tmp_path: Path, brief):
    """Секция «Ранние разработки и команды» печатается при непустом
    `contest_cards`; `developer_line` не тянет хвост «— не найдено»."""
    main_card = _make_card(title="Основная карточка")
    contest_card = _make_card(
        title="Конкурсная карточка",
        developer_org="ООО «Ромашка»",
        developer_team="не найдено",
    )
    report = _make_report([main_card], contest_cards=[contest_card], brief=brief)
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    assert "Ранние разработки и команды" in text
    assert "Конкурсная карточка" in text
    assert "ООО «Ромашка» — не найдено" not in text


def test_ограничения_и_риски_и_господдержка_печатаются(tmp_path: Path, brief):
    """Раздел ограничений, риски, господдержка и построчная сверка требований —
    все опциональные блоки карточки отражаются в тексте документа."""
    card = _make_card(
        risks=["Долгий цикл поставки присадки"],
        gov_support=[GovSupportMatch(instrument="Грант ФСИ", why="Подходит по отрасли")],
        requirement_checks=[RequirementCheck(requirement="Работа при 340 °C", status="confirmed")],
    )
    report = _make_report(
        [card],
        brief=brief,
        sources_unavailable=["cyberleninka"],
    )
    view = _make_view(report, is_draft=True)

    export_docx(view, tmp_path / "report-1-2026-08-27.docx")

    text = _full_text(tmp_path / "report-1-2026-08-27.docx")
    assert "Ограничения этого поиска" in text
    assert "Не ответили источники" in text
    assert "Долгий цикл поставки присадки" in text
    assert "Грант ФСИ" in text
    assert "Работа при 340 °C" in text


def test_создаёт_недостающий_каталог(tmp_path: Path, brief):
    """`out_path` в несуществующем подкаталоге — каталог создаётся."""
    report = _make_report([_make_card()], brief=brief)
    view = _make_view(report, is_draft=True)
    out_path = tmp_path / "vault" / "report-1-2026-08-27.docx"

    result = export_docx(view, out_path)

    assert result.exists()
