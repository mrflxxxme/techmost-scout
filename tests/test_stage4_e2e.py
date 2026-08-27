"""Сквозной тест этапа 4 (Э4-И8, PLAN.md волна 3, SPEC §6 Э4-1…Э4-3).

Один сценарий на фикстуре сохранённого прогона, без сети и без браузера
(инвариант №7): вердикты по всем карточкам запроса (один `reject`, один
`rework`, остальные `approve`) → статус запроса `validated` → страница отчёта
показывает подпись эксперта, прячет отклонённую карточку и помечает карточку
на доработку → `scout export --pdf` (страница Chromium подменена, реальный
`render_print_html` не тронут) и `--docx` дают ту же выборку карточек и ту же
подпись → демо-воспроизведение того же запроса (`app.pipeline.demo.play` с
подменённым ожиданием) не меняет `requests.cost_rub`.

Сценарий собирается один раз на модуль (`scope="module"`) — фикстура сама
проводит все шаги (сидирование базы, вердикты через `TestClient`, CLI-выгрузку,
демо-показ) и складывает готовые артефакты (HTML страницы, HTML печати, текст
DOCX) в `_Scenario`; тестовые функции только проверяют факты по уже собранным
артефактам — так пять с лишним проверок не пересобирают сценарий заново.

Приёмы подмен — с образцов задач волны 0–2:
- база — временный файл SQLite, подмена `app.db.session_scope`/`app.db.init_db`
  (форма `tests/test_web_review.py`, `tests/test_cli_export.py`);
- вердикты — `POST /cards/{id}/verdict` через `TestClient` (`tests/test_web_review.py`);
- `scout export` — `typer.testing.CliRunner`, `app.cli.get_settings` подменён
  на временный каталог (`tests/test_cli_export.py`); в PDF-выгрузке подменяется
  не вся функция `export_pdf`, а только страница Chromium (`page.pdf`) —
  `render_print_html` при этом настоящий, ровно так же, как в
  `tests/test_export_pdf.py` (`_FakePage`), а не мок всей функции;
- демо — `app.pipeline.demo.build_playback`/`play` вызываются напрямую на
  строках `pipeline_logs`, ожидание между шагами подменено (`tests/test_demo_playback.py`).
  Явный большой `request_id` — тот же приём, что в докстроке `_make_request`
  `tests/test_web_demo.py`: брокер прогресса (`app.web.progress`) общий
  синглтон на процесс, маленькие номера заняты другими сквозными файлами.
"""

from __future__ import annotations

import asyncio
import re
from collections.abc import AsyncIterator, Iterator
from contextlib import asynccontextmanager
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from typer.testing import CliRunner

import app.cli as cli
import app.cli_export as cli_export
from app import main
from app.export.markdown import REVIEWED_FOOTER
from app.export.pdf import export_pdf as real_export_pdf
from app.models import Base, RequestStatus
from app.models import PipelineLog as PipelineLogRow
from app.models import Request as RequestRow
from app.models import SolutionCard as SolutionCardRow
from app.pipeline import demo
from app.schemas import Brief, Evidence, Report, ReportSummary, SolutionCard
from app.web import progress

_AUTH = (main.settings.web_auth_user, main.settings.web_auth_password.get_secret_value())

#: Большой номер запроса — брокер прогресса общий синглтон на процесс, а
#: маленькие номера заняты другими сквозными файлами (`tests/test_web_demo.py`
#: докстрока `_make_request`).
_REQUEST_ID = 971_001

_REVIEWER = "Мария Петрова"

#: Заголовки карточек: два одобренных, один отклонённый, один на доработку —
#: ровно расклад из постановки задачи («один reject, один rework, остальные
#: approve»).
_APPROVE_TITLE_1 = "Катализатор с присадкой молибдена"
_APPROVE_TITLE_2 = "Катализатор с добавкой лантана"
_REJECT_TITLE = "Экспериментальная присадка на основе вольфрама"
_REWORK_TITLE = "Присадка на основе платины требует доработки"

#: Строка подписи (РЕШ-197) — вырезается из трёх собранных видов по общему
#: началу `REVIEWED_FOOTER`, дата и имя вычисляются по факту вердикта, а не
#: захардкожены — иначе тест был бы завязан на дату запуска.
_SIGNATURE_PATTERN = re.compile(r"Подготовлено системой ТехМост Скаут, проверено: .+?, \d{2}\.\d{2}\.\d{4}")


def _card(title: str, **overrides: object) -> SolutionCard:
    fields: dict[str, object] = {
        "title": title,
        "essence": "Присадка продлевает межрегенерационный цикл катализатора гидроочистки.",
        "applicability": "Закрывает задачу продления срока службы катализатора.",
        "next_step": "Запросить протокол испытаний.",
        "evidence": [Evidence(claim="Испытан на действующей установке", url="https://example.org/proof")],
    }
    fields.update(overrides)
    return SolutionCard(**fields)


def _brief() -> Brief:
    return Brief(
        industry="нефтепереработка",
        task_summary="Катализатор гидроочистки теряет активность быстрее ожидаемого.",
    )


def _report(cards: list[SolutionCard]) -> Report:
    return Report(
        request_id=_REQUEST_ID,
        brief=_brief(),
        cards=cards,
        contest_cards=[],
        summary=ReportSummary(recommendation="Начать с варианта №1."),
        generated_at=datetime(2026, 8, 27, 9, 0),
    )


def _docx_full_text(path: Path) -> str:
    """Весь текст DOCX одной строкой — абзацы и ячейки таблиц (приём `tests/test_export_docx.py`)."""
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _signature_in(text: str) -> str | None:
    match = _SIGNATURE_PATTERN.search(text)
    return match.group(0) if match else None


@dataclass
class _Scenario:
    """Готовые артефакты одного сквозного прогона — тесты только их читают."""

    status_after_verdicts: str
    report_html: str
    print_html: str
    docx_text: str
    cost_before: float
    cost_after: float
    sleep_calls: list[float]


class _FakePage:
    """Заглушка страницы Chromium: `set_content` запоминает HTML, `pdf` не запускает браузер."""

    def __init__(self) -> None:
        self.html: str | None = None

    async def set_content(self, html: str) -> None:
        self.html = html

    async def pdf(self, **_kwargs: object) -> bytes:
        return b"%PDF-1.4 fake"


@pytest.fixture(scope="module")
def scenario(tmp_path_factory: pytest.TempPathFactory) -> Iterator[_Scenario]:
    """Проводит весь сценарий один раз и отдаёт собранные артефакты.

    Модульный `pytest.MonkeyPatch()` вместо фикстуры `monkeypatch` — она
    функциональная и модульной фикстуре недоступна (её область шире).
    """
    mp = pytest.MonkeyPatch()
    tmp_path = tmp_path_factory.mktemp("stage4-e2e")

    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'stage4.db'}", future=True)

    async def _prepare() -> None:
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)

    asyncio.run(_prepare())
    session_maker = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)

    @asynccontextmanager
    async def factory() -> AsyncIterator[AsyncSession]:
        async with session_maker() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    mp.setattr("app.db.session_scope", factory)

    async def fake_init_db() -> None:
        return None

    mp.setattr("app.db.init_db", fake_init_db)
    mp.setattr(cli, "get_settings", lambda: SimpleNamespace(out_dir=tmp_path / "out"))

    captured_page = _FakePage()

    async def fake_export_pdf(view, out_path):
        # Подменяется только страница Chromium — сборка HTML идёт настоящим
        # `render_print_html`, чтобы проверять реальную выборку карточек, а не
        # заглушку (РЕШ-196 п. 12: браузер в наборе не поднимается).
        return await real_export_pdf(view, out_path, page=captured_page)

    mp.setattr(cli_export, "export_pdf", fake_export_pdf)

    cost_before = 7.42
    cards = [
        _card(_APPROVE_TITLE_1),
        _card(_APPROVE_TITLE_2),
        _card(_REJECT_TITLE),
        _card(_REWORK_TITLE),
    ]
    report = _report(cards)

    async def _seed() -> list[int]:
        async with factory() as session:
            row = RequestRow(
                id=_REQUEST_ID,
                raw_text="катализатор теряет активность раньше срока",
                status=RequestStatus.DRAFTED,
                cost_rub=cost_before,
                report_json=report.model_dump_json(),
            )
            session.add(row)

            card_rows = []
            for position, card in enumerate(cards):
                card_row = SolutionCardRow(
                    request_id=_REQUEST_ID,
                    position=position,
                    is_contest=False,
                    title=card.title,
                    essence=card.essence,
                    maturity_stage=card.maturity_stage.value,
                    applicability=card.applicability,
                    next_step=card.next_step,
                    evidence=[item.model_dump(mode="json") for item in card.evidence],
                    confidence=card.confidence.value,
                )
                session.add(card_row)
                card_rows.append(card_row)

            base = datetime(2026, 8, 27, 9, 0, 0, tzinfo=UTC)
            for stage, offset, duration in [("planner", 0, 10), ("search", 10, 20), ("scorer", 30, 70)]:
                started = base + timedelta(seconds=offset)
                session.add(
                    PipelineLogRow(
                        request_id=_REQUEST_ID,
                        stage=stage,
                        started_at=started,
                        finished_at=started + timedelta(seconds=duration),
                        sources_hit={"crossref": 3} if stage == "search" else {},
                    )
                )

            await session.flush()
            return [int(card_row.id) for card_row in card_rows]

    card_ids = asyncio.run(_seed())

    # ── Вердикты: один reject, один rework, остальные approve ───────────────
    decisions = [("approve", 5), ("approve", 4), ("reject", 2), ("rework", 3)]
    client = TestClient(main.app)
    with client:
        for card_id, (decision, rating) in zip(card_ids, decisions, strict=True):
            response = client.post(
                f"/cards/{card_id}/verdict",
                data={"decision": decision, "rating": str(rating), "comment": "", "reviewer_name": _REVIEWER},
                auth=_AUTH,
            )
            assert response.status_code == 200, response.text

        report_response = client.get(f"/requests/{_REQUEST_ID}/report", auth=_AUTH)
        assert report_response.status_code == 200
        report_html = report_response.text

    async def _status() -> str:
        async with factory() as session:
            found = (await session.execute(select(RequestRow).where(RequestRow.id == _REQUEST_ID))).scalars().one()
            return str(found.status)

    status_after_verdicts = asyncio.run(_status())

    # ── Выгрузка PDF и DOCX той же командой, которой пользуется оператор ────
    runner = CliRunner()
    pdf_result = runner.invoke(cli.app, ["export", str(_REQUEST_ID), "--pdf"])
    assert pdf_result.exit_code == 0, pdf_result.output
    print_html = captured_page.html
    assert print_html is not None

    docx_result = runner.invoke(cli.app, ["export", str(_REQUEST_ID), "--docx"])
    assert docx_result.exit_code == 0, docx_result.output
    docx_path = tmp_path / "out" / cli_export.export_filename(_REQUEST_ID, "docx")
    docx_text = _docx_full_text(docx_path)

    # ── Демо-воспроизведение того же запроса не должно менять cost_rub ──────
    sleep_calls: list[float] = []

    async def fake_sleep(seconds: float) -> None:
        sleep_calls.append(seconds)

    async def _playback() -> float:
        async with factory() as session:
            logs = (
                (
                    await session.execute(
                        select(PipelineLogRow)
                        .where(PipelineLogRow.request_id == _REQUEST_ID)
                        .order_by(PipelineLogRow.started_at)
                    )
                )
                .scalars()
                .all()
            )
        steps = demo.build_playback(logs, total_seconds=5.0)
        await demo.play(_REQUEST_ID, steps, broker=progress.get_broker(), sleep=fake_sleep)

        async with factory() as session:
            found = (await session.execute(select(RequestRow).where(RequestRow.id == _REQUEST_ID))).scalars().one()
            return float(found.cost_rub)

    cost_after = asyncio.run(_playback())

    yield _Scenario(
        status_after_verdicts=status_after_verdicts,
        report_html=report_html,
        print_html=print_html,
        docx_text=docx_text,
        cost_before=cost_before,
        cost_after=cost_after,
        sleep_calls=sleep_calls,
    )

    mp.undo()
    asyncio.run(engine.dispose())


def test_вердикты_по_всем_карточкам_переводят_запрос_в_validated(scenario: _Scenario) -> None:
    """После вердиктов по всем четырём карточкам статус запроса — `validated`."""
    assert scenario.status_after_verdicts == RequestStatus.VALIDATED.value


def test_страница_отчёта_показывает_нужную_выборку_и_подпись(scenario: _Scenario) -> None:
    """Клиентская страница: одобренные видны, отклонённая — нет, доработка помечена, подпись есть."""
    html = scenario.report_html
    assert _APPROVE_TITLE_1 in html
    assert _APPROVE_TITLE_2 in html
    assert _REJECT_TITLE not in html
    assert _REWORK_TITLE in html
    assert "На доработку" in html
    assert "ЧЕРНОВИК" not in html
    assert _signature_in(html) is not None


def test_pdf_печать_даёт_ту_же_выборку_и_подпись(scenario: _Scenario) -> None:
    """Печатный HTML (тот же, что уходит в `page.set_content`) — та же выборка карточек."""
    html = scenario.print_html
    assert _APPROVE_TITLE_1 in html
    assert _APPROVE_TITLE_2 in html
    assert _REJECT_TITLE not in html
    assert _REWORK_TITLE in html
    assert "На доработку" in html
    assert "ЧЕРНОВИК" not in html
    assert _signature_in(html) is not None


def test_docx_даёт_ту_же_выборку_и_подпись(scenario: _Scenario) -> None:
    """Текст DOCX — та же выборка карточек, что на странице и в PDF."""
    text = scenario.docx_text
    assert _APPROVE_TITLE_1 in text
    assert _APPROVE_TITLE_2 in text
    assert _REJECT_TITLE not in text
    assert _REWORK_TITLE in text
    assert "ЧЕРНОВИК" not in text
    assert _signature_in(text) is not None


def test_подпись_дословно_совпадает_во_всех_трёх_видах(scenario: _Scenario) -> None:
    """Строка РЕШ-197 — ровно одна и та же (символ в символ) на странице, в PDF и в DOCX."""
    page_signature = _signature_in(scenario.report_html)
    print_signature = _signature_in(scenario.print_html)
    docx_signature = _signature_in(scenario.docx_text)

    assert page_signature is not None
    assert print_signature is not None
    assert docx_signature is not None
    assert page_signature == print_signature == docx_signature
    assert page_signature.startswith(REVIEWED_FOOTER.format(name="", date="").split(",", 1)[0])
    assert _REVIEWER in page_signature


def test_демо_воспроизведение_того_же_запроса_не_меняет_cost_rub(scenario: _Scenario) -> None:
    """Показ демо-прогона не тратит ни рубля: `cost_rub` до и после совпадают."""
    assert scenario.cost_after == pytest.approx(scenario.cost_before)
    # Показ действительно прошёл через все шаги (не упал в пустой список) —
    # значит проверка выше не обходит стороной работающий код воспроизведения.
    assert scenario.sleep_calls
