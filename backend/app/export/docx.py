"""Выгрузка отчёта в DOCX (SPEC Э4-2, часть 2).

Данные приходят готовым `ExportView` из `export.view.build_export_view` —
единственного сборщика (РЕШ-196 п. 8): свой разбор `report_json`, свой фильтр
reject-карточек или повторное наложение правок этот файл не ведёт. Человеческие
подписи разделов и строки берутся из `export.markdown` — той же точкой
правды, что и markdown-экспорт, страница отчёта и печатный PDF-шаблон, чтобы
все три вида не разошлись формулировками.

Обложки нет (РЕШ-202): DOCX — рабочий документ для внутренней передачи, а не
питч-дек, этим он и отличается от PDF.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Pt

from app.export.analogs import (
    ANALOG_FOOTNOTES_TITLE,
    ANALOG_SECTION_TITLE,
    AnalogSections,
    analog_evidence,
    analog_fields,
    footnote_line,
)
from app.export.markdown import (
    COMPARISON_HEADERS_RU,
    COMPARISON_TABLE_HIDDEN_NOTICE,
    CONFIDENCE_RU,
    CONSTRAINTS_UNRESOLVED_NOTICE,
    CONTEST_SECTION_HEADING,
    MATURITY_RU,
    SUMMARY_NOT_PRINTABLE_NO_TABLE_NOTICE,
    SUMMARY_NOT_PRINTABLE_NOTICE,
    developer_line,
    limitations_lines,
)
from app.export.view import ExportView, caveats_printable, recommendation_printable
from app.schemas import Report, SolutionCard
from app.verdicts import signature_line

#: Пометка черновика — дословно существующая фраза (`report.html`, `markdown.py`).
DRAFT_NOTICE = "ЧЕРНОВИК — не проверено экспертом"

#: Шрифт документа — Segoe UI (SPEC Э4-2). DOCX хранит только имя шрифта;
#: системный запасной шрифт подставляет читалка, если Segoe UI не установлен.
_FONT_NAME = "Segoe UI"


def export_docx(view: ExportView, out_path: Path) -> Path:
    """Собрать DOCX-выгрузку `view` и сохранить её по пути `out_path`.

    Имя файла и каталог (`out/report-<id>-<дата>.docx`, РЕШ-196 п. 9) задаёт
    вызывающая сторона (команда `scout export`, Э4-И6) — сама функция пишет
    туда, куда указано, и создаёт недостающие родительские каталоги.
    """
    document = Document()
    _set_default_font(document)

    if view.is_draft:
        # Колонтитул — на каждой странице (П-4 аудита Э4): строка статуса
        # ниже видна только на первой странице документа.
        _add_draft_header(document)

    document.add_heading("Подбор научно-технических решений", level=0)
    document.add_paragraph(_status_line(view))

    # Источник — `view.client_report`, а не `view.report` (РЕШ-213): признак
    # видимости для клиента (reject или пустой `evidence`, РЕШ-204) считает
    # один раз `verdicts.visible_report`, эта функция своего фильтра по
    # `evidence` не ведёт — иначе строка таблицы сравнения могла бы остаться
    # напечатанной для карточки, которой в документе уже нет (БЛОК-6
    # финального аудита). Считаем до раздела «Рекомендация»: пополевой гейт
    # РЕШ-211 (`recommendation_printable`/`caveats_printable`) сравнивает
    # число реально печатаемых карточек с числом карточек до фильтра.
    report = view.client_report
    printable_cards = report.cards
    printable_contest_cards = report.contest_cards
    printed_total = len(printable_cards) + len(printable_contest_cards)
    printable_recommendation = recommendation_printable(view, printed_total)
    printable_caveats = caveats_printable(view, printed_total)
    # РЕШ-214: заглушка вместо рекомендации выбирается по факту «таблица
    # напечатана» (`bool(comparison_rows)` уже после фильтра), а не по причине
    # её отсутствия — `view.comparison_table_hidden` называет только одну
    # конкретную причину и остаётся отдельным решением («печатать ли на месте
    # таблицы честную строку», см. `_add_comparison_table`).
    table_printed = bool(report.summary.comparison_rows)

    _add_task_section(document, report)
    _add_recommendation_section(document, report, printable_recommendation, table_printed)
    _add_comparison_table(document, report, view.comparison_table_hidden)

    document.add_heading("Варианты решений", level=1)
    if not printable_cards:
        document.add_paragraph("Пригодных вариантов не найдено.")
    for index, card in enumerate(printable_cards, start=1):
        _add_card(document, card, index)

    if printable_contest_cards:
        document.add_heading(CONTEST_SECTION_HEADING.removeprefix("## "), level=1)
        for index, card in enumerate(printable_contest_cards, start=1):
            _add_card(document, card, index)

    _add_analog_section(document, view.analog_sections)

    _add_limitations_section(document, report, printable_caveats)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)
    return out_path


def _add_draft_header(document: DocumentObject) -> None:
    """Пометка черновика в колонтитуле — видна на каждой странице (П-4 аудита).

    Строка статуса в теле документа (`_status_line`) остаётся только на
    первой странице; колонтитул печатается на всех. У проверенного документа
    эта функция не вызывается вовсе — колонтитул остаётся пустым.
    """
    header = document.sections[0].header
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    paragraph.text = DRAFT_NOTICE


def _set_default_font(document: DocumentObject) -> None:
    """Базовый шрифт документа — Segoe UI, без внешних подключений."""
    style = document.styles["Normal"]
    style.font.name = _FONT_NAME
    style.font.size = Pt(11)


def _status_line(view: ExportView) -> str:
    """Пометка черновика или подпись эксперта — единственная строка статуса.

    Тот же порядок решения, что у страницы отчёта (`web/routes/report.py`):
    подпись строится, только если запрос не черновик и действующий вердикт
    уже есть; иначе печатается пометка черновика (инвариант №2).
    """
    if not view.is_draft and view.validated_at is not None:
        return signature_line(view.reviewer_name, view.validated_at)
    return DRAFT_NOTICE


def _add_task_section(document: DocumentObject, report: Report) -> None:
    """Раздел «Задача» — отрасль, суть, требования, желаемая готовность."""
    brief = report.brief
    document.add_heading("Задача", level=1)
    document.add_paragraph(f"Отрасль: {brief.industry}")
    document.add_paragraph(brief.task_summary)
    if report.constraints_unresolved:
        document.add_paragraph(CONSTRAINTS_UNRESOLVED_NOTICE)
    if brief.constraints:
        document.add_paragraph("Требования и ограничения: " + "; ".join(brief.constraints))
    if brief.target_maturity:
        document.add_paragraph(f"Требуемая готовность: {MATURITY_RU[brief.target_maturity]}")


def _add_recommendation_section(document: DocumentObject, report: Report, printable: bool, table_printed: bool) -> None:
    """Раздел «Рекомендация» — заглушка вместо текста модели (РЕШ-206, РЕШ-211),
    пока хоть одна карточка скрыта из клиентского вида и эксперт рекомендацию
    не правил.

    `table_printed` (РЕШ-212, РЕШ-214) — факт «таблица сравнения напечатана в
    этом же документе» (`bool(report.summary.comparison_rows)` уже после
    фильтра `verdicts.visible_report`), а не причина её отсутствия: обычная
    заглушка ссылается на таблицу сравнения как на ориентир и годится только
    когда таблица реально есть — иначе отсылать клиента к разделу, которого в
    документе нет, нельзя. Раньше здесь смотрели на `ExportView.
    comparison_table_hidden` — один конкретный повод спрятать таблицу, а не на
    сам факт; аудит нашёл два случая мимо него (сводка без строк сравнения при
    скрытой карточке и отчёт, где отклонены все карточки основной секции) —
    таблицы не было, а заглушка всё равно на неё ссылалась.
    """
    document.add_heading("Рекомендация", level=1)
    if printable:
        document.add_paragraph(report.summary.recommendation)
    elif table_printed:
        document.add_paragraph(SUMMARY_NOT_PRINTABLE_NOTICE)
    else:
        document.add_paragraph(SUMMARY_NOT_PRINTABLE_NO_TABLE_NOTICE)


def _add_comparison_table(document: DocumentObject, report: Report, table_hidden: bool) -> None:
    """Таблица сравнения вариантов — пусто, если `comparison_rows` не заполнены.

    Печатается всегда, вне зависимости от гейта рекомендации: строки уже
    отфильтрованы вместе с карточками (Э4-П1), домысливать в них нечего.

    `table_hidden` (РЕШ-212, РЕШ-213) — таблицу спрятал `verdicts.
    visible_report`, потому что построчное соответствие строк карточкам не
    сошлось (исходные строки были, но после фильтра — reject или пустой
    `evidence`, тот же признак, что у `ExportView.client_report` — их число не
    совпало с числом карточек). Раздел в этом случае не исчезает молча — на
    месте таблицы печатается честная строка `COMPARISON_TABLE_HIDDEN_NOTICE`.
    Пустая сводка модели без единой строки с самого начала (`table_hidden=
    False`) раздел по-прежнему не печатает вовсе — там скрывать нечего.
    """
    rows = report.summary.comparison_rows
    if not rows:
        if table_hidden:
            document.add_heading("Сравнение вариантов", level=1)
            document.add_paragraph(COMPARISON_TABLE_HIDDEN_NOTICE)
        return
    document.add_heading("Сравнение вариантов", level=1)
    keys = list(rows[0].keys())
    table = document.add_table(rows=1, cols=len(keys))
    table.style = "Table Grid"
    for cell, key in zip(table.rows[0].cells, keys, strict=True):
        cell.text = COMPARISON_HEADERS_RU.get(key, key)
    for row in rows:
        cells = table.add_row().cells
        for cell, key in zip(cells, keys, strict=True):
            cell.text = str(row.get(key, ""))


def _add_card(document: DocumentObject, card: SolutionCard, index: int) -> None:
    """Одна карточка: суть, разработчик, применимость, риски, следующий шаг,
    подтверждения. Тело общее для основной и конкурсной секций, как и в
    markdown-экспорте, — проверка инварианта №1 одна на обе секции.
    """
    document.add_heading(f"{index}. {card.title}", level=2)

    # Без единого подтверждения карточка не имеет права заявлять зрелость —
    # машинная проверка инварианта №1, повторяет ветку `markdown._render_card`.
    if card.evidence:
        maturity = MATURITY_RU.get(card.maturity_stage, card.maturity_stage.value)
        trl = f" (TRL ≈ {card.trl_estimate})" if card.trl_estimate else ""
        document.add_paragraph(f"Стадия зрелости: {maturity}{trl}")

    document.add_paragraph(card.essence)
    document.add_paragraph(f"Разработчик: {developer_line(card)}")
    if card.contact_leads:
        document.add_paragraph("Как выйти на связь: " + "; ".join(card.contact_leads))

    document.add_paragraph(f"Применимость к задаче: {card.applicability}")

    if card.requirement_checks:
        document.add_paragraph("Сверка с требованиями задачи:")
        for check in card.requirement_checks:
            status_text = "подтверждено источником" if check.status == "confirmed" else "не подтверждено источником"
            note = f" — {check.note}" if check.note else ""
            document.add_paragraph(f"{check.requirement}: {status_text}{note}", style="List Bullet")

    if card.risks:
        document.add_paragraph("Риски:")
        for risk in card.risks:
            document.add_paragraph(risk, style="List Bullet")

    if card.gov_support:
        document.add_paragraph("Господдержка:")
        for item in card.gov_support:
            limit = f" — {item.limit}" if item.limit else ""
            url = f" ({item.url})" if item.url else ""
            document.add_paragraph(f"{item.instrument}{limit}: {item.why}{url}", style="List Bullet")

    document.add_paragraph(f"Следующий шаг: {card.next_step}")

    confidence = CONFIDENCE_RU.get(card.confidence, card.confidence.value)
    note = f" — {card.confidence_note}" if card.confidence_note else ""
    document.add_paragraph(f"Уверенность: {confidence}{note}")

    document.add_paragraph("Чем подтверждено:")
    if card.evidence:
        for item in card.evidence:
            document.add_paragraph(f"{item.claim} — {item.url}", style="List Bullet")
    else:
        document.add_paragraph("Нельзя показывать клиенту: нет подтверждений")

    if card.unsourced_claims:
        document.add_paragraph("Не подтверждено проверкой:")
        for claim in card.unsourced_claims:
            document.add_paragraph(claim, style="List Bullet")


def _add_analog_section(document: DocumentObject, sections: AnalogSections) -> None:
    """Секция «Что делают в мире» (Э6-4): полные карточки и сноски; пусто — секции нет.

    Что видно, решил `export.analogs.analog_sections` (подтверждения и страна
    режима) — этот файл своего фильтра не ведёт.
    """
    if not sections:
        return
    document.add_heading(ANALOG_SECTION_TITLE, level=1)
    for index, card in enumerate(sections.main, start=1):
        document.add_heading(f"{index}. {card.title}", level=2)
        document.add_paragraph(card.essence)
        for label, value in analog_fields(card):
            document.add_paragraph(f"{label}: {value}")
        document.add_paragraph("Чем подтверждено:")
        for claim, url in analog_evidence(card):
            document.add_paragraph(f"{claim} — {url}", style="List Bullet")
    if sections.footnotes:
        document.add_paragraph(f"{ANALOG_FOOTNOTES_TITLE}:")
        for footnote in sections.footnotes:
            document.add_paragraph(footnote_line(footnote), style="List Bullet")


def _add_limitations_section(document: DocumentObject, report: Report, printable: bool) -> None:
    """Раздел «Ограничения этого поиска» — тем же текстом, что markdown и веб.

    `printable` — результат `caveats_printable` (РЕШ-206, РЕШ-211), не общего
    гейта: `printable=False` убирает только свободную прозу модели — строка о
    недоступных источниках и машинные пометки (платный предел поиска, отказ
    интейка/рефайна/сводки) остаются в любом случае (инвариант №5, РЕШ-210).
    """
    lines = limitations_lines(report, include_caveats=printable)
    if not lines:
        return
    document.add_heading("Ограничения этого поиска", level=1)
    for line in lines:
        if line:
            document.add_paragraph(line.removeprefix("- "))
